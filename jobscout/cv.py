"""Read a CV from disk and return plain text.

Supports .pdf, .docx, .md, .txt and .tex. PDF and DOCX need optional
dependencies; the error message says which one to install.
"""

from __future__ import annotations

import re
from pathlib import Path

MAX_CHARS = 60_000


class CVError(Exception):
    """Raised when a CV cannot be read."""


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise CVError(
            "Reading PDFs needs pypdf. Install it with: pip install 'jobscout[pdf]'"
        ) from exc
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(path: Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise CVError(
            "Reading .docx needs python-docx. Install it with: pip install 'jobscout[docx]'"
        ) from exc
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _strip_latex(text: str) -> str:
    text = re.sub(r"(?m)%.*$", "", text)
    text = re.sub(r"\\(?:textbf|textit|emph|texttt|url|href)\{", "", text)
    text = re.sub(r"\\[a-zA-Z@]+\*?(?:\[[^\]]*\])?", " ", text)
    return text.replace("{", " ").replace("}", " ").replace("\\&", "&")


READERS = {
    ".pdf": _read_pdf,
    ".docx": _read_docx,
}

PLAIN_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".tex"}


def read_cv(path: str | Path) -> str:
    """Return the text of the CV at ``path``.

    Raises CVError if the file is missing, empty, or in a format that
    cannot be read.
    """
    p = Path(path).expanduser()
    if not p.is_file():
        raise CVError(f"No file at {p}")

    suffix = p.suffix.lower()
    if suffix in READERS:
        text = READERS[suffix](p)
    elif suffix in PLAIN_SUFFIXES:
        text = p.read_text(encoding="utf-8", errors="replace")
        if suffix == ".tex":
            text = _strip_latex(text)
    else:
        raise CVError(
            f"Cannot read '{suffix}' files. Supported: "
            + ", ".join(sorted(set(READERS) | PLAIN_SUFFIXES))
        )

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 200:
        raise CVError(
            f"{p.name} yielded only {len(text)} characters of text. "
            "If it is a scanned PDF, export a text version first."
        )
    return text[:MAX_CHARS]
